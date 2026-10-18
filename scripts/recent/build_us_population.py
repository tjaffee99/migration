"""Build data/processed/us_population_annual.csv  (year,population,basis,source_id)

Census counts for the census years 1790-1890 (intercensal years left out on purpose: the page
interpolates), Census Bureau July 1 national estimates for 1900-2025, and the Vintage 2025
short-term projection for July 1, 2026.

Raw inputs (fetch with scripts/recent/fetch_recent.sh):
  data/raw/census_wp56/appendixA.xlsx            WP56 Appendix A, US row of Tables A-15..A-26
  data/raw/census_popest/popclockest.txt         1900-1999 national estimates (used 1900-1989)
  data/raw/census_popest/us-est90int-07.csv      1990-2000 intercensal by age (All Age, July 1)
  data/raw/census_popest/us-est00int-tot.csv     2000-2010 intercensal, monthly totals
  data/raw/census_popest/nst-est2020int-pop.xlsx 2010-2020 intercensal (released Nov 2024)
  data/raw/census_popest/NST-EST2025-ALLDATA.csv Vintage 2025 (released Jan 2026)
  data/raw/census_popest/NA-EST2025-POP.xlsx     Vintage 2025 monthly national, to Dec 1, 2026
Run: python3 -I scripts/recent/build_us_population.py
"""
import csv, os, re
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
RAW = os.path.join(ROOT, 'data', 'raw')
OUT = os.path.join(ROOT, 'data', 'processed', 'us_population_annual.csv')
PE = os.path.join(RAW, 'census_popest')


def wp56_census_counts():
    wb = openpyxl.load_workbook(os.path.join(RAW, 'census_wp56', 'appendixA.xlsx'), read_only=True,
                                data_only=True)
    out = {}
    for sn in wb.sheetnames:
        m = re.match(r'(\d{4}) Table A-(\d+)', sn)
        if not m or int(m.group(1)) > 1890 or sn == '1890 Table A-16':   # A-16 = 1890 general enumeration only
            continue
        y = int(m.group(1))
        for r in wb[sn].iter_rows(values_only=True):
            cells = [c for c in r if c is not None and str(c).strip() != '']
            if cells and isinstance(cells[0], str) and cells[0].strip().startswith('United States') \
                    and len(cells) > 1 and isinstance(cells[1], (int, float)):
                out[y] = int(cells[1])
                break
        else:
            raise SystemExit(f'no United States row in {sn}')
    return out


def popclock():
    out = {}
    for line in open(os.path.join(PE, 'popclockest.txt'), encoding='latin-1'):
        m = re.match(r'\s*July 1, (\d{4})\s+([\d,]+)\s', line)
        if m:
            out[int(m.group(1))] = int(m.group(2).replace(',', ''))
    return out


def int90():
    out = {}
    for r in csv.reader(open(os.path.join(PE, 'us-est90int-07.csv'), encoding='latin-1')):
        if len(r) >= 3 and r[0].startswith('July 1,') and r[1].strip() == 'All Age':
            out[int(r[0][-4:])] = int(r[2])
    return out


def int00():
    out = {}
    for r in csv.DictReader(open(os.path.join(PE, 'us-est00int-tot.csv'))):
        if r['MONTH'] == '7':
            out[int(r['YEAR'])] = int(r['TOT_POP'])
    return out


def int00_table1():
    """BOTH SEXES row of us-est00int-01.csv (Table 1, resident population), July 1 2000-2009."""
    rows = list(csv.reader(open(os.path.join(PE, 'us-est00int-01.csv'), encoding='latin-1')))
    yrs = next(r for r in rows if len(r) > 3 and r[2] == '2000')
    both = next(r for r in rows if r and r[0] == 'BOTH SEXES')
    return {int(y): int(v.replace(',', '')) for y, v in zip(yrs, both) if y.isdigit()}


def int10():
    ws = openpyxl.load_workbook(os.path.join(PE, 'nst-est2020int-pop.xlsx'), data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hdr = next(r for r in rows if r and r[1] == 'April 1, 2010 Estimates Base')
    yrs = next(r for r in rows if r and r[2] == 2010)
    us = next(r for r in rows if r and r[0] == 'United States')
    return {int(y): int(v) for y, v in zip(yrs, us) if isinstance(y, int)}


def v2025():
    r = next(x for x in csv.DictReader(open(os.path.join(PE, 'NST-EST2025-ALLDATA.csv'),
                                             encoding='latin-1')) if x['NAME'] == 'United States')
    return {y: int(r[f'POPESTIMATE{y}']) for y in range(2020, 2026)}


def v2025_monthly():
    ws = openpyxl.load_workbook(os.path.join(PE, 'NA-EST2025-POP.xlsx'), data_only=True).active
    out, year = {}, None
    for r in ws.iter_rows(values_only=True):
        a = r[0]
        if isinstance(a, int) or (isinstance(a, str) and re.match(r'^\d{4}( \[1\])?$', a)):
            year = int(str(a)[:4])
        elif isinstance(a, str) and a.startswith('.') and year:
            out[(year, a.strip('.').strip())] = dict(resident=r[1], resident_af=r[2], civ=r[3],
                                                    civ_noninst=r[4], household=r[5])
    return out


def main():
    rows = []
    cc = wp56_census_counts()
    for y in range(1790, 1891, 10):
        rows.append((y, cc[y], 'census_count', 'census_wp56_appA'))
    pc = popclock()
    for y in range(1900, 1990):
        basis = ('estimate_jul1_resident_excl_ak_hi' if y < 1940 else
                 'estimate_jul1_resident_plus_af_overseas_excl_ak_hi' if y < 1950 else
                 'estimate_jul1_resident_plus_af_overseas' if y < 1980 else
                 'estimate_jul1_resident')
        if y < 1930:
            basis += '_rounded_1000'
        rows.append((y, pc[y], basis, 'census_popest_1900_1999'))
    i90, i00, i10, v25 = int90(), int00(), int10(), v2025()
    for y in range(1990, 2000):
        rows.append((y, i90[y], 'intercensal_jul1_resident', 'census_popest_int1990_2000'))
    t1 = int00_table1()
    assert all(t1[y] == i00[y] for y in range(2000, 2010)), 'us-est00int-tot != Table 1'
    for y in range(2000, 2010):
        rows.append((y, i00[y], 'intercensal_jul1_resident', 'census_popest_int2000_2010'))
    for y in range(2010, 2020):
        rows.append((y, i10[y], 'intercensal_jul1_resident', 'census_popest_int2010_2020'))
    mo = v2025_monthly()
    for y in range(2020, 2026):
        assert mo[(y, 'July 1')]['resident'] == v25[y], (y, mo[(y, 'July 1')], v25[y])
        rows.append((y, v25[y], 'postcensal_v2025_jul1_resident', 'census_popest_v2025'))
    rows.append((2026, int(mo[(2026, 'July 1')]['resident']), 'projection_v2025_jul1_resident',
                 'census_popest_v2025_monthly'))
    rows.sort()
    with open(OUT, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['year', 'population', 'basis', 'source_id'])
        w.writerows(rows)
    print(f'wrote {OUT}: {len(rows)} rows, {rows[0][0]}-{rows[-1][0]}')
    # cross-checks between overlapping series (printed, quoted in recent_NOTES.md)
    print('checks:')
    print(f'  1990 popclockest {pc[1990]:,} vs intercensal {i90[1990]:,} ({pc[1990] - i90[1990]:+,})')
    print(f'  1999 popclockest {pc[1999]:,} vs intercensal {i90[1999]:,} ({pc[1999] - i90[1999]:+,})')
    print(f'  1999 int1990-2000 {i90[1999]:,} -> 2000 int2000-2010 {i00[2000]:,} '
          f'({i00[2000] - i90[1999]:+,})')
    print(f'  2010 int2000-2010 {i00[2010]:,} vs int2010-2020 {i10[2010]:,}')
    print(f'  2026 July 1 projection {mo[(2026, "July 1")]["resident"]:,}; '
          f'2025->2026 change {mo[(2026, "July 1")]["resident"] - v25[2025]:+,}')


if __name__ == '__main__':
    main()
